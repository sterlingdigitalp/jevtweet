"""Minimal understood-subset xlsx sheet reading (header row + data rows)."""

from __future__ import annotations

from pathlib import Path


def read_sheet(path: str | Path, sheet: str, *, max_rows: int = 5000) -> list[dict]:
    """Return data rows as header-keyed dicts; blank header cells end the row shape."""
    import openpyxl

    path = Path(path)
    if path.stat().st_size > 25_000_000:
        raise ValueError("Workbook exceeds 25 MB limit")
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        if sheet not in workbook.sheetnames:
            raise ValueError(f"Unknown sheet: {sheet}")
        rows = list(workbook[sheet].iter_rows(values_only=True))
    finally:
        workbook.close()
    if not rows:
        return []
    width = next((i for i, value in enumerate(rows[0]) if value in (None, "")), len(rows[0]))
    header = [str(rows[0][i]).strip() for i in range(width)]
    if len(header) != len(set(header)):
        raise ValueError("Duplicate header labels are ambiguous")
    out = []
    for values in rows[1:]:
        if all(value in (None, "") for value in values[:width]):
            continue
        out.append({key: values[i] for i, key in enumerate(header)})
    if len(out) > max_rows:
        raise ValueError(f"Sheet exceeds {max_rows} row cap")
    return out
